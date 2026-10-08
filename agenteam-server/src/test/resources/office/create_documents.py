"""生成可重复的中文办公样本，只写入测试容器的 outputs 目录。"""

from pathlib import Path
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font

output = Path("/workspace/outputs")
document = Document()
document.styles["Normal"].font.name = "Noto Sans CJK SC"
document.styles["Normal"].font.size = Pt(11)
document.sections[0].header.paragraphs[0].text = "AgenTeam 文档验证"
document.sections[0].footer.paragraphs[0].text = "测试资料，不包含用户数据"
document.add_heading("中文办公文档验证", 0)
document.add_paragraph("生成、读取、编辑和下载使用同一份文件。")
table = document.add_table(rows=1, cols=3)
for cell, value in zip(table.rows[0].cells, ["项目", "数量", "说明"]):
    cell.text = value
for index in range(45):
    cells = table.add_row().cells
    cells[0].text = f"文档项目 {index + 1}"
    cells[1].text = str(index + 1)
    cells[2].text = "验证中文表格跨页后内容仍完整。"
document.save(output / "sample.docx")

workbook = Workbook()
sheet = workbook.active
sheet.title = "计算"
sheet.append(["项目", "数量"])
sheet.append(["第一项", 2])
sheet.append(["第二项", 3])
sheet.append(["合计", "=SUM(B2:B3)"])
sheet.column_dimensions["A"].width = 22
sheet.column_dimensions["B"].width = 16
for row in sheet:
    for cell in row:
        cell.font = Font(name="Noto Sans CJK SC", size=11)
workbook.create_sheet("说明")["A1"] = "公式必须重新计算后核对数值"
workbook.save(output / "sample.xlsx")
