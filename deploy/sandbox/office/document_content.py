"""按文档类型检查并提取内容，供模型阅读和附件解析复用。"""

import hashlib
import json
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import ZipFile, BadZipFile

from defusedxml import ElementTree

MAX_EXPANDED = 100 * 1024 * 1024


class DocumentFailure(ValueError):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code


def validate_office(path, kind, budget=None, depth=0):
    if budget is None:
        budget = [0, 0]
    if depth > 2:
        raise DocumentFailure("FILE_ACTIVE_CONTENT", "文件嵌套层数超过允许范围")
    required = {"docx": "word/document.xml", "xlsx": "xl/workbook.xml", "pptx": "ppt/presentation.xml"}
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            names = set()
            for entry in entries:
                name = entry.filename
                parts = PurePosixPath(name).parts
                if name in names or name.startswith("/") or "\\" in name or ".." in parts:
                    raise DocumentFailure("FILE_TYPE_INVALID", "文件包含重复或无效路径")
                names.add(name)
                budget[0] += entry.file_size
                budget[1] += 1
                if budget[1] > 10000 or budget[0] > MAX_EXPANDED or entry.file_size > MAX_EXPANDED:
                    raise DocumentFailure("FILE_EXPANDED_TOO_LARGE", "文件展开后的内容过大")
                if entry.flag_bits & 1:
                    raise DocumentFailure("FILE_ENCRYPTED", "文件已加密")
                if "vbaproject" in name.lower() or "/activex/" in name.lower():
                    raise DocumentFailure("FILE_ACTIVE_CONTENT", "文件包含宏或嵌入程序")
                if "/embeddings/" in name.lower():
                    if kind == "pptx" and name.lower().endswith(".xlsx"):
                        validate_office(BytesIO(archive.read(entry)), "xlsx", budget, depth + 1)
                    elif not entry.is_dir():
                        raise DocumentFailure("FILE_ACTIVE_CONTENT", "文件包含不支持的嵌入程序")
            if required[kind] not in names or "[Content_Types].xml" not in names:
                raise DocumentFailure("FILE_TYPE_INVALID", "文件格式与扩展名不符")
            types = ElementTree.fromstring(archive.read("[Content_Types].xml"))
            expected = {
                "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
            }
            if not any(item.attrib.get("ContentType") == expected[kind] for item in types):
                raise DocumentFailure("FILE_TYPE_INVALID", "文件不是受支持的普通办公文档")
    except (BadZipFile, KeyError) as failure:
        raise DocumentFailure("FILE_TYPE_INVALID", "文件损坏或格式不正确") from failure


def sections(path, kind):
    if kind in {"docx", "xlsx", "pptx"}:
        validate_office(path, kind)
    if kind == "docx":
        from docx import Document
        document = Document(path)
        for index, paragraph in enumerate(document.paragraphs, 1):
            if paragraph.text.strip():
                yield None, f"段落 {index}", paragraph.text
        for index, table in enumerate(document.tables, 1):
            for row_index, row in enumerate(table.rows, 1):
                yield None, f"表格 {index} 第 {row_index} 行", " | ".join(cell.text for cell in row.cells)
    elif kind == "xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            cells = 0
            for sheet in workbook:
                sheet.reset_dimensions()
                yield None, f"工作表 {sheet.title}", f"工作表：{sheet.title}"
                for row in sheet.iter_rows():
                    cells += len(row)
                    if cells > 1000000:
                        raise DocumentFailure("FILE_EXPANDED_TOO_LARGE", "表格单元格数量超过限制")
                    values = [f"{cell.coordinate}={cell.value}" for cell in row if cell.value is not None]
                    if values:
                        yield None, f"工作表 {sheet.title}", " | ".join(values)
        finally:
            workbook.close()
    elif kind == "pptx":
        from pptx import Presentation
        presentation = Presentation(path)
        for number, slide in enumerate(presentation.slides, 1):
            yield number, f"幻灯片 {number}", f"第 {number} 页"
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    yield number, f"幻灯片 {number}", shape.text
                if shape.has_table:
                    for row in shape.table.rows:
                        yield number, f"幻灯片 {number} 表格", " | ".join(cell.text for cell in row.cells)
    elif kind == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise DocumentFailure("FILE_ENCRYPTED", "文件已加密")
        if len(reader.pages) > 2000:
            raise DocumentFailure("FILE_EXPANDED_TOO_LARGE", "文档页数超过限制")
        for number, page in enumerate(reader.pages, 1):
            yield number, f"第 {number} 页", page.extract_text() or ""
    elif kind in {"txt", "md"}:
        with open(path, encoding="utf-8-sig") as source:
            for line in source:
                yield None, None, line.rstrip("\r\n")
    else:
        raise DocumentFailure("FILE_TYPE_INVALID", "不支持此文档格式")


def extract(path, kind, output):
    count = 0
    total = 0
    pages = None
    with open(output, "w", encoding="utf-8") as target:
        for page, section, text in sections(path, kind):
            if not text.strip():
                continue
            if any(ord(char) < 32 and char not in "\n\r\t" for char in text):
                raise DocumentFailure("FILE_TYPE_INVALID", "文件包含无法读取的控制字符")
            if page:
                pages = max(page, pages or 0)
            offset = 0
            while offset < len(text):
                part = text[offset:offset + 800]
                count += 1
                block = {"ordinal": count, "page": page, "section": section, "text": part,
                         "sha256": hashlib.sha256(part.encode("utf-8")).hexdigest(), "overlapCharacters": 0}
                encoded = json.dumps(block, ensure_ascii=False) + "\n"
                total += len(encoded.encode("utf-8"))
                if total > MAX_EXPANDED or count > 200000:
                    raise DocumentFailure("FILE_EXPANDED_TOO_LARGE", "提取的文档内容超过限制")
                target.write(encoded)
                offset += 800
    if not count:
        raise DocumentFailure("FILE_NO_TEXT", "文件没有可读取的文字")
    return {"success": True, "errorCode": None, "chunkCount": count, "pageCount": pages, "csv": None}
